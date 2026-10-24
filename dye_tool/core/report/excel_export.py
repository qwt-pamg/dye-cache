# -*- coding: utf-8 -*-
"""P5：Excel 多 sheet 导出 —— 把**一次 run** 的产出汇总成一个 ``.xlsx`` 工作簿。

设计约定
========

1. **唯一输入是** :class:`core.pipeline.RunResult` 与完整配置 ``cfg``。本模块
   **不重新拟合、不重算 k 与 Φ_Δ、不读原始数据**（唯一的读盘例外是 run 目录里的
   ``run_params.yaml`` 与染料档案 ``history.csv``），因此导出的数与前一刻算出的数
   **同源**——这正是"导出的数和算出的数不一样"这类事故的根治办法。
2. **IO 层绝不打断 GUI**（本项目铁律）：:func:`export_workbook` 顶层 try/except
   兜底，任何失败（未装 openpyxl、目录只读、磁盘满、run_result 残缺…）都只记
   WARNING 并返回 ``None``，**绝不把异常抛给调用方**。
3. **单表失败不废整本**：每张表单独构建；某张表数据缺失时写"空表 + 提示行"，
   其余表照常生成。
4. **数值保持数值**：数值列写的是 ``int``/``float``（不是字符串），只按
   ``decimals`` 设置 Excel 的**数字格式**（显示位数），单元格里存的是**全精度**
   数值——下游在 Excel 里可以直接参与计算。
5. **core 层无 GUI 依赖**：本模块不 import matplotlib / PySide6。

表结构（表名 = ``cfg["excel"]["sheets"]`` 的键，关掉的表不进工作簿）
===================================================================

* ``summary``  每染料 × 每组一行：时间点数、k ± σ、R²、n、A0、Φ_Δ ± σ、参比、
  截断、状态、备注、对照判定、run 目录
* ``qc``       逐文件一行：t、数据点数、无效点数、异常点数、探针窗有效点、
  是否饱和、A0、A0 池化、是否参与拟合、剔除原因、QC 状态
* ``kinetics`` 拟合逐点明细：t、横轴 x、A、A0、A/A0、lnR、是否参与拟合、
  是否被截断、说明
* ``dyes``     染料档案 / 跨 run 历史：k ± σ、R²、Φ ± σ、run 时间、来源
* ``params``   本次 run 的参数快照（``run_params.yaml`` 摊平成「参数 / 值」两列）

两个口径说明（如实写在这里，避免下游误读）
==========================================

* **「异常点数」的口径**：``Scan`` 上**没有**记录去尖峰实际剔除的点数
  （:func:`core.qc.despike` 只把异常点置 NaN，不落盘计数），因此本模块把
  「异常点数」定义为**探针窗口内**的非有限点数——它正是会直接影响 A(t) 与 k 的
  那部分坏点；「无效点数」则是整条光谱的非有限点数（含仪器缺测）。
  若需要"despike 实际剔除数"，需在 ``core.qc`` 侧把计数写回 ``Scan``
  （属 core 数据模型改动，不在 P5 边界内）。
* **空表形态**：表头仍在第 1 行（列契约不因缺数据而失效），提示行
  ``（无数据：<原因>）`` 写在第 2 行 A 列——这样"冻结首行/自动筛选"在空表上
  依然是正确的。
"""
from __future__ import annotations

import logging
import math
import os
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import numpy as np

from ..config import section_defaults

log = logging.getLogger("dye_tool.excel")

# 优化 #7（V2.x 审查 W4）：openpyxl 占启动 import ≈0.77 s，而 Excel 导出是**用户
# 主动动作**。改为首次导出时才导入（模块级缓存）。GUI 启动完全不付这份成本；缺
# 依赖时仍在真正导出那一刻给出可读提示（行为契约不变）。
_OPENPYXL_PROBED, _OPENPYXL_OK, _IMPORT_ERR = False, False, None
Workbook = None              # type: ignore[assignment]
Font = None                  # type: ignore[assignment]
get_column_letter = None     # type: ignore[assignment]


def _ensure_openpyxl() -> bool:
    """首次调用时 import openpyxl 并缓存结果；返回是否可用。"""
    global Workbook, Font, get_column_letter
    global _OPENPYXL_PROBED, _OPENPYXL_OK, _IMPORT_ERR
    if _OPENPYXL_PROBED:
        return bool(_OPENPYXL_OK)
    _OPENPYXL_PROBED = True
    try:
        from openpyxl import Workbook as _WB
        from openpyxl.styles import Font as _Font
        from openpyxl.utils import get_column_letter as _gcl
        Workbook, Font, get_column_letter = _WB, _Font, _gcl
        _OPENPYXL_OK, _IMPORT_ERR = True, None
    except Exception as err:                        # noqa: BLE001 —— 可选依赖
        _OPENPYXL_OK, _IMPORT_ERR = False, err
    return bool(_OPENPYXL_OK)

__all__ = ["export_workbook", "SHEET_ORDER",
           "write_queue_summary_workbook", "write_statistics_workbook"]

#: 工作表顺序（与 ``DEFAULT_CFG["excel"]["sheets"]`` 的键一一对应）
SHEET_ORDER = ("summary", "qc", "kinetics", "dyes", "params")

#: 缺数据提示行的前缀（中文表头/正文由 openpyxl 以 UTF-8 写入，Excel 正常显示）
NOTE_PREFIX = "（无数据："
_NOTE_FALLBACK = "本次 run 没有可写出的行"
_DEFAULT_FILENAME = "dye_tool_结果汇总.xlsx"

# 这两个字体对象原来在 import 期构造（需要 openpyxl）→ 现在由 `_styles()` 懒建。
_BOLD = None
_WARN_FONT = None


def _styles():                                     # noqa: ANN202 —— 返回 (Font|None, Font|None)
    """首次导出时构造粗体/警示字体（openpyxl 未装则为 None）。"""
    global _BOLD, _WARN_FONT
    if _BOLD is None and _ensure_openpyxl():
        _BOLD = Font(bold=True)
        _WARN_FONT = Font(bold=True, color="FFC00000")
    return _BOLD, _WARN_FONT


# ═══════════════════════════ 取值 / 类型小工具 ═══════════════════════════
def _attr(obj: Any, name: str, default: Any = None) -> Any:
    """同时兼容 ``RunResult`` 与 dict（缺字段返回 default，永不抛异常）。"""
    if isinstance(obj, dict):
        return obj.get(name, default)
    return getattr(obj, name, default)


def _as_path(v: Any) -> Optional[Path]:
    """任意值 → Path；空/不可用 → None（注意 ``Path("")`` 是当前目录，必须挡住）。"""
    if v is None:
        return None
    try:
        s = str(v).strip()
    except Exception:                                   # noqa: BLE001
        return None
    return Path(s) if s else None


def _dict(v: Any) -> dict:
    """只接受 dict（字段被写成别的类型时按"空"处理，绝不迭代出垃圾行）。"""
    return v if isinstance(v, dict) else {}


def _seq(v: Any) -> list:
    """只接受序列（**字符串不算**——否则会被逐字符迭代出垃圾行）。"""
    if isinstance(v, (list, tuple)):
        return list(v)
    return []


def _fnum(v: Any) -> Optional[float]:
    """尽力转 float；失败或非有限 → None（Excel 里留空，不写 nan/inf）。"""
    if v is None or isinstance(v, (bool, np.bool_)):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _nm(v: Any) -> Any:
    """数值列取值：数值原样（``int`` 保持 ``int``），其余 → None。"""
    if isinstance(v, (bool, np.bool_)) or v is None:
        return None
    if isinstance(v, (int, np.integer)):
        return int(v)
    return _fnum(v)


def _int_or_none(v: Any) -> Optional[int]:
    f = _fnum(v)
    return int(round(f)) if f is not None else None


def _cell(v: Any) -> Any:
    """任意值 → openpyxl 认得的单元格值（数值保持数值；NaN/±Inf → 空单元格）。"""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        f = float(v)
        return f if math.isfinite(f) else None
    if isinstance(v, str):
        return v
    if isinstance(v, Path):
        return str(v)
    return str(v)


def _text(v: Any) -> str:
    """任意值 → 单元格文本（None → 空串）。"""
    return "" if v is None else str(v)


def _num_fmt(v: float, decimals: int) -> str:
    """按 ``decimals`` 生成 Excel 数字格式；极小/极大值用科学计数法。

    只改**显示格式**，单元格里存的仍是全精度数值（下游可直接计算）。
    """
    try:
        d = max(0, min(15, int(decimals)))
    except (TypeError, ValueError):
        d = 6
    frac = "0" * d
    a = abs(float(v))
    if a == 0.0 or 1e-4 <= a < 1e7:
        return f"0.{frac}" if d else "0"
    return f"0.{frac}E+00" if d else "0E+00"


def _disp_len(s: str) -> int:
    """列宽用的显示宽度：CJK 全角字符算 2 个字符位。"""
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
               for ch in str(s))


def _stamp(run_id: str) -> str:
    """run 目录名/run_id → ``"2026-10-14 22:21:04"``；认不出则原样返回。

    M8（评审记录）：``pipeline`` 传进来的 run_id 恒为 ``run_dir.name``，
    而 V1.0 遗留目录是 ``run_20261014-222104``（**带 ``run_`` 前缀**）。
    原实现只顾着取前 15 个字符去 ``strptime``，于是前缀把时间戳挤掉一位 →
    **必然失配 → 恒返回原始 run_id**，dyes 表的"run 时间"列永远显示
    ``run_20261014-222104`` 这种机器名。

    修法：先把 ``run_`` 前缀剥掉，再在字符串里**找**时间戳（而不是假定它在开头）
    ——这样新形态 ``20261014-222104``、旧形态 ``run_20261014-222104``、
    以及带序号后缀的 ``20261014-222104_2`` 都能正确格式化。
    """
    s = str(run_id or "").strip()
    if s.startswith("run_"):
        s = s[4:]
    import re
    m = re.match(r"(\d{8}-\d{6})", s)
    if not m:
        return str(run_id or "")
    try:
        return datetime.strptime(m.group(1), "%Y%m%d-%H%M%S").strftime(
            "%Y-%m-%d %H:%M:%S")
    except ValueError:                                  # pragma: no cover —— 正则已限形
        return str(run_id or "")


def _run_cfg(run_result: Any, cfg: dict) -> dict:
    """取"本次 run 真正用的那份配置"（含 ``_dye_session``），兜底用调用方给的。"""
    rc = _attr(run_result, "cfg")
    if isinstance(rc, dict) and rc:
        return rc
    return cfg or {}


def _out_root(cfg: dict) -> Optional[Path]:
    """配置里的输出根目录（V2.5 D10：唯一解析在 core.paths.resolve_out_root）。"""
    try:
        from ..paths import resolve_out_root
        return resolve_out_root(cfg)
    except Exception:                                   # noqa: BLE001
        return None


def _dye_name(cfg: dict, group: str) -> str:
    """组名 → 染料名（**与 run 输出目录同源**：会话 > 注册表 > auto）。"""
    try:
        from .dyes import dye_for_group
        return str(dye_for_group(cfg, group))
    except Exception as err:                            # noqa: BLE001
        log.debug("染料名解析失败(组=%s): %s", group, err)
        return str(group or "")


def _n_points(reps: dict) -> int:
    """各重复的时间点总数（数组不能直接求真值，逐个 size 累加）。"""
    n = 0
    for r in (reps or {}).values():
        t = (r or {}).get("t")
        if t is None:
            continue
        try:
            n += int(np.asarray(t).size)
        except Exception:                               # noqa: BLE001
            continue
    return n


def _fit_note(fit: Any, ctrl: Any) -> tuple[str, str]:
    """``(备注, 对照判定)``——与 ``results.csv`` 的 note 组法保持一致。"""
    msg = _text(getattr(fit, "message", "") or "")
    judge = why = ""
    if isinstance(ctrl, (list, tuple)) and len(ctrl) >= 2:
        judge, why = _text(ctrl[0]), _text(ctrl[1])
    note = (msg + (f" | {why}" if why else "")).strip(" |")
    return note, judge


# ═══════════════════════════ 工作表写出内核 ═══════════════════════════
#: Excel 工作表的硬上限（行）。第 1 行是表头，所以**数据行**最多 XLSX_MAX_ROWS-1。
#: L19（评审记录）：openpyxl 自己**不拦**这个上限——实测行号推过限后 ``save()``
#: 不报错、回读 ``max_row=1048577``，而那个 xlsx 在 Excel 里打开会弹
#: "发现不可读取的内容，是否修复"。触发面是超大 run 的 kinetics 明细
#: （行数 = Σ重复 × 时间点）。所以必须自己截断 + 明确告警。
XLSX_MAX_ROWS = 1_048_576


def _write_table(ws, headers, rows, *, decimals: int, freeze: bool,
                 autofilter: bool, width_max: float, note: Optional[str] = None,
                 trunc_note: Optional[str] = None) -> int:
    """把一张表写进 ``ws``：表头 + 数据行（无数据则写提示行）+ 格式。

    表头始终在第 1 行（列契约不因缺数据而失效），因此 ``freeze_panes="A2"`` 与
    ``auto_filter`` 在空表上也保持正确语义。

    Returns:
        实际写入的**数据行数**（被截断时为上限值）。
    """
    ws.append([_cell(h) for h in headers])
    if _BOLD is not None:
        for c in ws[1]:
            c.font = _BOLD
    data = list(rows or [])
    # L19：超上限就截断，并在末尾追加一行醒目说明（绝不产出会让 Excel 报"修复"的文件）。
    #
    # G10（第二轮审查）：这里原来是 `limit = XLSX_MAX_ROWS - 1`，**差一行**——
    # 表头 1 行 + limit 行 = 恰好顶格 1,048,576；再加那行"已截断"说明就是
    # **第 1,048,577 行，超上限**，正是 L19 要防的形态（Excel 会弹"发现不可读取的
    # 内容，是否修复"；若 openpyxl 改为抛错，本函数在每表 try 之外，整本导出会中止）。
    # 现在 `- 2`：表头 + 数据 + 说明行 = 1,048,576，刚好顶格不超。
    # 原注释"说明行会落到 limit+2 ≤ 上限之内"的算术也是错的（limit+2 = MAX+1）。
    limit = XLSX_MAX_ROWS - 2
    truncated = len(data) > limit
    if truncated:
        data = data[:limit]
        log.warning("Excel 表「%s」行数超过上限，已截断到 %d 行（原 %d 行）",
                    getattr(ws, "title", "?"), limit, len(rows))
    # T9（评审记录）：原先这里写完之后再 `ws.iter_rows(min_row=2)` **整表遍历一遍**
    # 只为了给 float 单元格设 number_format —— 那等于把所有单元格对象化两次
    # （上限 10⁵ 行时是几百万次属性访问）。现在**边写边记**"哪一列是浮点列"，
    # 写完只回填少数几列（通常 ≤20 列），语义完全等价、遍历量降两个数量级。
    float_cols: set = set()
    if data:
        for r in data:
            vals = [_cell(x) for x in r]
            ws.append(vals)
            for ci, v in enumerate(vals, start=1):
                if isinstance(v, float):
                    float_cols.add(ci)
        if truncated:
            ws.append([NOTE_PREFIX + _text(
                trunc_note or f"数据过多，仅保留前 {limit} 行（共 {len(rows)} 行）") + "）"])
            if _WARN_FONT is not None:
                ws.cell(row=ws.max_row, column=1).font = _WARN_FONT
    else:
        ws.append([NOTE_PREFIX + _text(note or _NOTE_FALLBACK) + "）"])
        if _WARN_FONT is not None:
            ws.cell(row=2, column=1).font = _WARN_FONT
    # 数值列：存全精度数值，只按 decimals 设显示格式（只碰浮点列）
    last_row = max(1, int(ws.max_row or 1))
    for ci in sorted(float_cols):
        for r in range(2, last_row + 1):
            c = ws.cell(row=r, column=ci)
            if isinstance(c.value, float):
                c.number_format = _num_fmt(c.value, decimals)
    if freeze:
        ws.freeze_panes = "A2"
    if autofilter:
        last_col = get_column_letter(max(1, int(ws.max_column or 1)))
        last_row = max(1, int(ws.max_row or 1))
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    _fit_widths(ws, width_max)
    return len(data)


def _fit_widths(ws, width_max: float) -> None:
    """按内容自适应列宽并**封顶** ``width_max``（<=0 表示不设，交给 Excel 默认）。"""
    try:
        cap = float(width_max)
    except (TypeError, ValueError):
        cap = 60.0
    if not math.isfinite(cap) or cap <= 0:
        return
    for idx, col in enumerate(ws.iter_cols(), start=1):
        best = 0
        for c in col:
            if c.value is None:
                continue
            best = max(best, _disp_len(_text(c.value)))
        if best <= 0:
            continue
        ws.column_dimensions[get_column_letter(idx)].width = min(cap, float(best + 2))


# ═══════════════════════════ summary：每染料 × 每组一行 ═══════════════════════════
_SUMMARY_HEAD = ["染料", "组", "角色", "重复数", "时间点数", "k (s⁻¹)", "k 标准误",
                 "R²", "n", "A0", "A_ex", "辐照度 (mW/cm²)", "Φ_Δ", "Φ_Δ 误差",
                 "参比", "截断", "状态", "备注", "对照判定", "run 目录"]


def _summary_rows(run_result, cfg, ex) -> tuple[list, list, Optional[str]]:
    results = _dict(_attr(run_result, "results"))
    phi = _dict(_attr(run_result, "phi"))
    controls = _dict(_attr(run_result, "controls"))
    run_dir = _attr(run_result, "run_dir")
    rcfg = _run_cfg(run_result, cfg)
    rows: list = []
    for g in sorted(results, key=str):
        e = _dict(results.get(g))
        f = e.get("pooled")
        p = _dict(phi.get(g))
        reps = _dict(e.get("reps"))
        refs = "、".join(str(q.get("ref")) for q in _seq(p.get("pairs"))
                         if isinstance(q, dict) and q.get("ref"))
        note, judge = _fit_note(f, controls.get(g))
        status = _text(getattr(f, "status", "") or ("FAILED" if f is None else ""))
        rows.append([
            _dye_name(rcfg, g), _text(g), _text(e.get("role")), len(reps),
            _n_points(reps), _nm(getattr(f, "k", None)),
            _nm(getattr(f, "k_err", None)), _nm(getattr(f, "r2", None)),
            _nm(getattr(f, "n", None)), _nm(e.get("A0_pooled")),
            _nm(e.get("a_ex_auto")), _nm(e.get("irradiance")),
            _nm(p.get("phi")), _nm(p.get("phi_err")),
            refs or _text(p.get("status")),
            "是" if bool(getattr(f, "truncated", False)) else "否",
            status, note, judge, _text(run_dir),
        ])
    return _SUMMARY_HEAD, rows, ("本次 run 没有可拟合的组（results 为空）"
                                 if not rows else None)


# ═══════════════════════════ qc：逐文件一行 ═══════════════════════════
_QC_HEAD = ["文件", "组", "重复", "角色", "t", "数据点数", "无效点数", "异常点数",
            "探针窗有效点", "是否饱和", "A0", "A0 池化", "是否参与拟合",
            "剔除原因", "QC 状态", "备注"]


#: 三处"回退默认值"的**唯一**来源（重复逻辑 #9，评审记录）。
#: 原先 probe/win/sat/r_min/r_max 的兜底常量在 `_qc_rows` 与 `_kinetics_rows`
#: 里各自硬编码了一份，与 ``core.config.DEFAULT_CFG`` 是**第三**份。现值恰好一致，
#: 于是没人发现；一旦 config 默认改动就会**静默分叉**（表里的数与会话里的 cfg 不同）。
#:
#: V2.0 δ 批（FS2）：改成**直接从 DEFAULT_CFG 取**（经 :func:`core.config.section_defaults`），
#: 不再手抄字面量——这样"改了默认值忘了同步这里"这件事在结构上不可能发生。
#: 名字保持兼容（测试与既有调用点都引用 `_FB_*`）。
_FB_PROBE = float(section_defaults("wavelength").get("probe", 379.0))
_FB_PROBE_WIN = float(section_defaults("wavelength").get("probe_window", 2.0))
_FB_R_MIN = float(section_defaults("kinetics").get("r_min", 0.30))
_FB_R_MAX = float(section_defaults("kinetics").get("r_max", 1.00))
_FB_R_MAX_TOL = float(section_defaults("kinetics").get("r_max_tol", 0.05))


def _qc_rows(run_result, cfg, ex) -> tuple[list, list, Optional[str]]:
    scans = _seq(_attr(run_result, "scans"))
    results = _dict(_attr(run_result, "results"))
    wl_cfg = _dict((cfg or {}).get("wavelength"))
    try:
        probe = float(wl_cfg.get("probe", _FB_PROBE))
    except (TypeError, ValueError):
        probe = _FB_PROBE
    try:
        win = abs(float(wl_cfg.get("probe_window", _FB_PROBE_WIN)))
    except (TypeError, ValueError):
        win = _FB_PROBE_WIN
    rows: list = []
    for s in scans:
        g = _text(getattr(s, "group", ""))
        rep = _text(getattr(s, "rep", ""))
        role = _text(getattr(s, "role", ""))
        excluded = _text(getattr(s, "excluded", ""))
        entry = _dict(results.get(g))
        reps = _dict(entry.get("reps"))
        repd = _dict(reps.get(rep))
        rfit = repd.get("fit")
        arr = getattr(s, "a", None)
        wl = getattr(s, "wl", None)
        try:
            a = np.asarray(arr, float) if arr is not None else None
        except Exception:                               # noqa: BLE001
            a = None
        try:
            w = np.asarray(wl, float) if wl is not None else None
        except Exception:                               # noqa: BLE001
            w = None
        n_total = n_bad = n_probe_bad = n_probe_ok = None
        if a is not None:
            fin = np.isfinite(a)
            n_total, n_bad = int(a.size), int(a.size - np.count_nonzero(fin))
            if w is not None and w.size == a.size:
                m = np.abs(w - probe) <= win
                n_probe_bad = int(np.count_nonzero(m & ~fin))
                n_probe_ok = int(np.count_nonzero(m & fin))
        # L23（评审记录）：这里原来**自己复刻了一遍** flag_saturation 的判据
        # （`nanmax(a) > saturation`）。它继承了一个"量纲幽灵"：`qc.saturation`
        # 默认 3.0 是**吸光度**阈值，而荧光仪导出的是**原始强度**（10²–10⁴），
        # 于是实测出现过"强度 5000、未被剔除的文件却显示 饱和=是 / QC状态=正常"
        # ——同一张表两列自相矛盾。判据的唯一真相是**上游 QC 的结论**
        # （`Scan.excluded` 里的"饱和"字样），这里只读它，不再自己算。
        hot = "饱和" in excluded
        if excluded:
            state = "剔除"
        elif (n_probe_bad or n_bad):
            state = "注意"
        else:
            state = "正常"
        notes = []
        if n_probe_bad:
            notes.append(f"探针窗口内 {n_probe_bad} 个无效点（已置 NaN，不参与 A(t)）")
        if n_bad and n_probe_bad is not None and n_bad > n_probe_bad:
            notes.append(f"谱内另有 {n_bad - n_probe_bad} 个无效点（探针窗之外，"
                         "不影响 k）")
        used = (not excluded and role != "background" and bool(repd)
                and (rfit is None or _text(getattr(rfit, "status", "")) != "FAILED"))
        rows.append([
            _text(getattr(getattr(s, "path", None), "name", "")), g, rep, role,
            _nm(getattr(s, "t", None)), n_total, n_bad, n_probe_bad, n_probe_ok,
            "是" if hot else "否", _nm(repd.get("A0")), _nm(entry.get("A0_pooled")),
            "是" if used else "否", excluded, state, "；".join(notes),
        ])
    return _QC_HEAD, rows, ("本次 run 没有文件（scans 为空）" if not rows else None)


# ═══════════════════════════ kinetics：拟合逐点明细 ═══════════════════════════
_KIN_HEAD = ["染料", "组", "重复", "t", "横轴 x", "A", "A0", "A/A0", "lnR",
             "是否参与拟合", "是否被截断", "说明"]


def _kinetics_rows(run_result, cfg, ex) -> tuple[list, list, Optional[str]]:
    """拟合逐点明细。

    评审记录「重复逻辑 #5」：本函数**手工复刻**了 ``core.kinetics.fit_kinetics``
    的选窗判据（docstring 自认"同一判据"）。两处各写一遍必然漂移，所以改成调用
    :func:`core.kinetics.window_mask` —— 那是从 ``fit_kinetics`` 里**抽出来的**
    同一份判据，拟合与导出共用。
    """
    from ..analysis.kinetics import window_mask, window_params

    results = _dict(_attr(run_result, "results"))
    rcfg = _run_cfg(run_result, cfg)
    kin = _dict((cfg or {}).get("kinetics"))
    # r_min/r_max 只用于**表格文案与提示**（判据本身在 window_mask 里）；
    # 常量集中在模块顶部 _FB_*（重复逻辑 #9），不再与 config.py 各写一份。
    try:
        r_min, r_max = window_params(kin)
    except (TypeError, ValueError):
        r_min, r_max = _FB_R_MIN, _FB_R_MAX + _FB_R_MAX_TOL
    dose = bool(kin.get("dose_normalize", False))
    rows: list = []
    for g in sorted(results, key=str):
        e = _dict(results.get(g))
        reps = _dict(e.get("reps"))
        dye = _dye_name(rcfg, g)
        try:
            irr = float(e.get("irradiance", 1.0))
        except (TypeError, ValueError):
            irr = 1.0
        for rep in sorted(reps, key=str):
            r = _dict(reps.get(rep))
            t = _arr(r.get("t"))
            A = _arr(r.get("A"))
            R = _arr(r.get("R"))
            if t is None or A is None:
                continue
            n = min(t.size, A.size)
            R = R if (R is not None and R.size >= n) else None
            fit = r.get("fit")
            status = _text(getattr(fit, "status", "") or "")
            truncated = bool(getattr(fit, "truncated", False))
            x = t * irr / 1000.0 if dose else t
            # 与 core.kinetics.fit_kinetics **同一份判据**（重复逻辑 #5）：
            # `window_mask` 是从 fit_kinetics 里抽出来的那支，两边共用，
            # 不再各写一遍（原先的本地复刻被评审记录点名为"必然漂移"）。
            mask = window_mask(t[:n], R[:n], kin, x[:n]) if R is not None \
                else np.zeros(n, dtype=bool)
            idx = [int(i) for i in np.flatnonzero(mask)]
            order = sorted(idx, key=lambda i: float(x[i]))
            n_used = 0
            fx = getattr(fit, "x", None)
            try:
                n_used = 0 if fx is None else int(np.asarray(fx).size)
            except Exception:                           # noqa: BLE001
                n_used = 0
            used = set(order[:n_used]) if n_used else set()
            idx_set = set(idx)
            a0 = _fnum(r.get("A0"))
            for i in range(n):
                ri = float(R[i]) if R is not None else float("nan")
                if i in used:
                    ok, cut, why = "是", "否", ""
                else:
                    ok = "否"
                    cut = "是" if (truncated and i in idx_set) else "否"
                    if not math.isfinite(float(A[i])):
                        why = "A 无效(NaN)"
                    elif R is None or not math.isfinite(ri):
                        why = "R 无效（未归一化或 A 低于 a_floor）"
                    elif not (r_min <= ri <= r_max):
                        why = f"R={ri:.3f} 超出拟合窗 [{r_min:.2f}, {r_max:.2f}]"
                    elif cut == "是":
                        why = f"被自动截断舍弃（保留最早 {n_used}/{len(idx)} 点）"
                    elif status == "FAILED":
                        why = "拟合失败（该点未被使用）"
                    else:
                        why = "未参与拟合"
                rows.append([
                    dye, _text(g), _text(rep), _nm(t[i]), _nm(x[i]), _nm(A[i]), a0,
                    _nm(ri), _nm(math.log(ri) if math.isfinite(ri) and ri > 0 else None),
                    ok, cut, why,
                ])
    return _KIN_HEAD, rows, ("本次 run 没有任何拟合点（results 为空）"
                             if not rows else None)


def _arr(v) -> Optional[np.ndarray]:
    """尽力转 float 数组；不可用 → None。"""
    if v is None:
        return None
    try:
        return np.asarray(v, float).ravel()
    except Exception:                                   # noqa: BLE001
        return None


# ═══════════════════════════ dyes：染料档案 / 跨 run 历史 ═══════════════════════════
_DYE_HEAD = ["染料", "来源", "run 时间", "run_id", "组", "k (s⁻¹)", "k 标准误", "R²",
             "n", "A0", "Φ_Δ", "Φ_Δ 误差", "状态", "备注", "文件数", "有效数",
             "run 目录"]


def _dyes_rows(run_result, cfg, ex) -> tuple[list, list, Optional[str]]:
    run_dir = _as_path(_attr(run_result, "run_dir"))
    results = _dict(_attr(run_result, "results"))
    phi = _dict(_attr(run_result, "phi"))
    scans = _seq(_attr(run_result, "scans"))
    rcfg = _run_cfg(run_result, cfg)
    rows: list = []
    seen: set = set()
    for d in _find_dye_dirs(run_dir, cfg):
        for h in _read_history(d):
            h = _dict(h)
            rid = _text(h.get("run_id", ""))
            rows.append(_dye_row_values(d.name, h, "history.csv"))
            seen.add((str(d.name), rid))
    # 本次 run 尚未登记（例如关了 dye.history）→ 现场补一行，并把来源写清楚，
    # 免得用户把"现值"误读成"档案里已有的一行"。
    if run_dir is not None and results:
        per: dict = {}
        for g in sorted(results, key=str):
            per.setdefault(_dye_name(rcfg, g), []).append(str(g))
        n_files = len(scans)
        n_used = sum(1 for s in scans if not getattr(s, "excluded", ""))
        for dye, groups in sorted(per.items()):
            if (dye, run_dir.name) in seen:
                continue
            try:
                from .dyes import history_row
                h = history_row(run_dir.name, run_dir, dye, groups, results, phi,
                                n_files, n_used)
            except Exception as err:                    # noqa: BLE001
                log.debug("本次 run 档案行组装失败(染料=%s): %s", dye, err)
                continue
            rows.append(_dye_row_values(dye, h, "本次 run（未登记档案）"))
    return _DYE_HEAD, rows, ("未找到任何染料档案（既没有 history.csv，"
                             "本次 run 也没有可登记的组）" if not rows else None)


def _dye_row_values(dye: str, h: dict, src: str) -> list:
    rid = _text(h.get("run_id", ""))
    note = _text(h.get("notes", ""))
    if _fnum(h.get("time_edited")):
        note = (note + "；手动调整时间点").strip("；")
    return [dye, src, _stamp(rid), rid, _text(h.get("groups", "")),
            _fnum(h.get("k")), _fnum(h.get("k_se")), _fnum(h.get("R2")),
            _int_or_none(h.get("n")), _fnum(h.get("A0")), _fnum(h.get("Phi")),
            _fnum(h.get("Phi_se")), _text(h.get("status", "")), note,
            _int_or_none(h.get("n_files")), _int_or_none(h.get("n_used")),
            _text(h.get("dir", ""))]


def _read_history(dye_dir: Path) -> list:
    try:
        from .dyes import read_history
        return list(read_history(dye_dir) or [])
    except Exception as err:                            # noqa: BLE001
        log.debug("读取染料档案失败(%s): %s", dye_dir, err)
        return []


def _dye_dirs_under(root: Any) -> list:
    """``root`` 下**确实带 history.csv** 的染料目录（其余一律不算，避免误报）。"""
    r = _as_path(root)
    if r is None or not r.is_dir():
        return []
    try:
        from .dyes import HISTORY_NAME, scan_runs
        info = scan_runs(r) or {}
    except Exception as err:                            # noqa: BLE001
        log.debug("扫描染料目录失败(%s): %s", r, err)
        return []
    out: list = []
    for ent in _seq(info.get("dyes")):
        ent = _dict(ent)
        cands = _seq(ent.get("dirs"))
        d0 = ent.get("dir")
        if d0 is not None:
            cands.append(d0)
        for c in cands:
            p = _as_path(c)
            if p is not None and (p / HISTORY_NAME).exists():
                out.append(p)
    return out


def _is_dye_dir(p: Any) -> Optional[Path]:
    """``p`` 自身是否就是一个染料档案目录（即含 ``history.csv``）。"""
    q = _as_path(p)
    if q is None or not q.is_dir():
        return None
    try:
        from .dyes import HISTORY_NAME
    except Exception:                                   # noqa: BLE001
        return None
    return q if (q / HISTORY_NAME).exists() else None


def _dedup_paths(dirs: Any) -> list:
    """按**解析后的绝对路径**去重，保留出现顺序。

    ``scan_runs`` 的 ``dirs`` 与 ``dir`` 可能指向同一个目录（多形态归档），
    所以扫描结果天然会重复。
    """
    out: list = []
    seen: set = set()
    for d in _seq(dirs):
        p = _as_path(d)
        if p is None:
            continue
        try:
            key = str(p.resolve()).lower()
        except OSError:
            key = str(p).lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(p)
    return out


def _find_dye_dirs(run_dir: Optional[Path], cfg: dict) -> list:
    """定位**跟本次 run 相关**的染料档案目录，按出现顺序去重。

    L22（评审记录）：这个函数的 docstring 一直写着"跟本次 run 相关"，但实现是
    "把 run 目录的两级父目录都丢给 ``_dye_dirs_under`` 去**扫描**"，而
    ``runs/<染料>/<run_id>`` 的两级父目录正是 ``runs/`` 本身 → 扫出来的是
    **整个 runs/ 下所有染料**。实测 ``runs/DyeA/run_1`` 返回 ``['DyeB','DyeA']``，
    于是 dyes 表会混进与本次 run 无关的染料。

    现在的取法（与 docstring 一致）：

    1. ``run_dir.parent``（``runs/<染料>/``）**自身**带档案 → 就是它；
    2. 否则说明本 run 落在 ``runs/_multi/<run_id>`` 这类容器下 →
       扫 ``run_dir.parent.parent``（``runs/``）下的相关染料；
    3. 都取不到才退回配置里的输出根。

    保留"扫容器"这一支是有意的：``_multi`` 里一次 run 可能牵涉多个染料，
    它们的档案都该进 dyes 表。
    """
    if run_dir is not None:
        own = _is_dye_dir(run_dir.parent)
        if own is not None:
            return [own]
        out = _dedup_paths(_dye_dirs_under(run_dir.parent.parent))
        if out:
            return out
    return _dedup_paths(_dye_dirs_under(_out_root(cfg)))


# ═══════════════════════════ params：参数快照摊平 ═══════════════════════════
_PARAM_HEAD = ["参数", "值"]


def _params_rows(run_result, cfg, ex) -> tuple[list, list, Optional[str]]:
    run_dir = _as_path(_attr(run_result, "run_dir"))
    data: Optional[dict] = None
    src = ""
    if run_dir is not None:
        p = run_dir / "run_params.yaml"
        if p.exists():
            try:
                import yaml
                loaded = yaml.safe_load(p.read_text(encoding="utf-8"))
                if isinstance(loaded, dict) and loaded:
                    data, src = loaded, "run_params.yaml"
            except Exception as err:                    # noqa: BLE001
                log.debug("读取 run_params.yaml 失败: %s", err)
    if data is None:
        rcfg = _run_cfg(run_result, cfg) or {}
        data = {k: v for k, v in rcfg.items() if not str(k).startswith("_")}
        src = "当前配置快照（run_params.yaml 缺失或不可读）"
    if not data:
        return _PARAM_HEAD, [], "未找到 run_params.yaml，也没有可用的配置快照"
    return _PARAM_HEAD, [["（来源）", src]] + _flatten(data), None


def _flatten(d: dict, prefix: str = "") -> list:
    """嵌套 dict → ``[(dotted_key, value)]``；list 拼成一行文本。"""
    out: list = []
    for k, v in (d or {}).items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            if v:
                out += _flatten(v, key + ".")
            else:
                out.append([key, "（空）"])
        elif isinstance(v, (list, tuple)):
            if v:
                out.append([key, "、".join(_scalar_text(x) for x in v)])
            else:
                out.append([key, "（空）"])
        else:
            out.append([key, _cell(v)])
    return out


def _scalar_text(v: Any) -> str:
    if isinstance(v, (dict, list, tuple)):
        try:
            import json
            return json.dumps(v, ensure_ascii=False)
        except Exception:                               # noqa: BLE001
            return str(v)
    return _text(v)


#: 表名 → 构建函数
_BUILDERS = {
    "summary": _summary_rows,
    "qc": _qc_rows,
    "kinetics": _kinetics_rows,
    "dyes": _dyes_rows,
    "params": _params_rows,
}


# ═══════════════════════════ 公开接口 ═══════════════════════════
def export_workbook(run_result, cfg: dict, out_path=None) -> Optional[Path]:
    """把一次 run 的产出汇总成一个多 sheet 的 ``.xlsx``。

    Args:
        run_result: :func:`core.pipeline.run_phi_pipeline` 的返回值
            （``RunResult``；也容忍 dict 与"字段残缺的替身"，缺失字段按空处理）。
        cfg: 完整配置 dict（表开关取 ``cfg["excel"]["sheets"]``，文件名取
            ``cfg["excel"]["filename"]``）。
        out_path: 输出路径；``None`` → ``run_result.run_dir / cfg["excel"]["filename"]``。

    Returns:
        写出的工作簿路径；**任何失败都返回 ``None``**（只记日志，绝不抛异常）。
        ``excel.enabled=false``、所有表都被关掉、未装 openpyxl、run_dir 与
        out_path 都缺、写盘失败——一律 ``None``。
    """
    try:
        ex = (cfg or {}).get("excel") or {}
        if not bool(ex.get("enabled", True)):
            log.info("Excel 导出已在配置里关闭（excel.enabled=false）→ 跳过")
            return None
        if not _ensure_openpyxl():
            log.warning("未安装 openpyxl，无法导出 Excel 汇总（pip install openpyxl）：%s",
                        _IMPORT_ERR)
            return None
        run_dir = _as_path(_attr(run_result, "run_dir"))
        if out_path is None:
            if run_dir is None:
                log.warning("既没有 run_dir 也没有 out_path → 无法确定 Excel 输出位置，"
                            "本次跳过导出")
                return None
            out_path = run_dir / _text(ex.get("filename") or _DEFAULT_FILENAME)
        path = Path(out_path)

        # L20（评审记录）：`ex.get("sheets")` 返回 None 时**不能**直接当成 {}——
        # 那会让"手搓的 cfg 没写 sheets 键"变成"所有表都关掉"，返回 None 并打一条
        # 误导的日志（与 `enabled` 默认 True 的取向也不对称）。真配置经 deep_merge
        # 永远带这个键，所以这条只影响 API 直调；但默认值必须两端一致。
        sheets = ex.get("sheets")
        if not isinstance(sheets, dict) or not sheets:
            if sheets is not None and not isinstance(sheets, dict):
                log.warning("cfg['excel']['sheets'] 不是映射（实际 %s）→ 按「全部生成」处理",
                            type(sheets).__name__)
            else:
                log.info("cfg['excel'] 未给出 sheets 开关 → 按「全部生成」处理")
            sheets = {name: True for name in SHEET_ORDER}
        unknown = [k for k in sheets if k not in SHEET_ORDER]
        if unknown:
            log.warning("excel.sheets 里有本版本不认识的表名（已忽略）: %s", unknown)

        try:
            decimals = max(0, min(15, int(ex.get("decimals", 6))))
        except (TypeError, ValueError):
            decimals = 6
        try:
            width_max = float(ex.get("column_width_max", 60))
        except (TypeError, ValueError):
            width_max = 60.0
        freeze = bool(ex.get("freeze_header", True))
        autofilter = bool(ex.get("autofilter", True))

        _styles()

        wb = Workbook()
        try:
            wb.remove(wb.active)            # 去掉默认的空表，只留配置里打开的
        except Exception:                   # noqa: BLE001
            pass
        built: list = []
        for name in SHEET_ORDER:
            if not bool(sheets.get(name, False)):
                continue
            try:
                headers, rows, note = _BUILDERS[name](run_result, cfg, ex)
            except Exception as err:        # noqa: BLE001 —— 单表失败不废整本
                log.warning("Excel 表 %s 构建失败，改写成空表: %s", name, err)
                headers, rows, note = ["说明"], [], f"{name} 表构建失败：{err}"
            ws = wb.create_sheet(title=name)
            _write_table(ws, headers, rows, decimals=decimals, freeze=freeze,
                         autofilter=autofilter, width_max=width_max, note=note)
            built.append(name)
        if not built:
            # L20：区分"用户显式全关"与"表都构建失败了"——原来一律说成前者，
            # 后者会让用户以为是自己配置的问题，其实是导出出错。
            effective = [n for n in SHEET_ORDER if bool(sheets.get(n, False))]
            if not effective:
                log.warning("excel.sheets 把所有工作表都关掉了 → 不生成空工作簿"
                            "（至少把 summary 设为 true）")
            else:
                log.warning("Excel 没有任何工作表构建成功（应为 %s）→ 不生成空工作簿",
                            effective)
            return None

        path.parent.mkdir(parents=True, exist_ok=True)
        # 原子写（与 results.atomic_write 同思路）：先写同目录临时文件再 os.replace，
        # 中途失败/被杀不会留下"半截 xlsx"（Excel 打不开的那种）。
        tmp = path.with_name(f".{path.name}.tmp{os.getpid()}")
        try:
            wb.save(tmp)
            os.replace(tmp, path)
        except BaseException:
            try:
                tmp.unlink()
            except OSError:
                pass
            raise
        # 落地校验：调用方会拿这个路径去打开文件夹/写日志，
        # **返回一个不存在的路径**比返回 None 更坏（UI 会打开一个空目录、
        # 日志会宣称导出成功）。所以这里必须确认文件真的在。
        if not path.exists():
            log.warning("Excel 保存后目标文件不存在（%s）——可能被外部移动/占用", path)
            return None
        log.info("已写出 Excel 汇总: %s（%d 张表：%s）", path, len(built),
                 "、".join(built))
        return path
    except Exception as err:                            # noqa: BLE001 —— IO 层铁律：不抛
        log.warning("Excel 导出失败（不影响分析结果与既有工件）: %s", err)
        import traceback
        log.debug("%s", traceback.format_exc())
        return None


# ═══════════════ V2.5：队列汇总 / 多 run 统计（FR-29.3 / FR-30.2） ═══════════════
def _xlsx_opts(cfg: dict) -> tuple[int, float, bool, bool]:
    """从 cfg 的 excel 节取显示选项（缺省与 export_workbook 同源口径）。"""
    ex = (cfg or {}).get("excel") or {}
    try:
        decimals = max(0, min(15, int(ex.get("decimals", 6))))
    except (TypeError, ValueError):
        decimals = 6
    try:
        width_max = float(ex.get("column_width_max", 60))
    except (TypeError, ValueError):
        width_max = 60.0
    return decimals, width_max, bool(ex.get("freeze_header", True)), \
        bool(ex.get("autofilter", True))


def _save_workbook_atomic(wb, out_path) -> Optional[Path]:
    """原子保存工作簿（临时文件 + os.replace）并做落地校验；失败返回 None。"""
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name("." + path.name + ".tmp" + str(os.getpid()))
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise
    if not path.exists():
        log.warning("工作簿保存后目标文件不存在（%s）——可能被外部移动/占用", path)
        return None
    return path


def write_queue_summary_workbook(rows, out_path, cfg=None) -> Optional[Path]:
    """写出批量队列汇总 xlsx（V2.5 FR-29.3：**一个目录一行**）。

    Args:
        rows: :func:`core.report.summary.collect_queue_summary` 的行 dict 列表。
        out_path: 输出路径（调用方用 summary.default_output_path 生成，避免覆盖）。
        cfg: 可选配置（只借用 excel 节的 decimals/列宽/冻结/筛选）。

    Returns:
        写出的路径；空行/未装 openpyxl/写盘失败 → None（只记日志，绝不抛）。
    """
    try:
        if not rows:
            log.info("队列汇总无数据行 → 不生成空文件")
            return None
        if not _ensure_openpyxl():
            log.warning("未安装 openpyxl，无法导出队列汇总（pip install openpyxl）：%s",
                        _IMPORT_ERR)
            return None
        from .summary import QUEUE_HEADERS
        decimals, width_max, freeze, autofilter = _xlsx_opts(cfg or {})
        _styles()
        wb = Workbook()
        try:
            wb.remove(wb.active)
        except Exception:                               # noqa: BLE001
            pass
        ws = wb.create_sheet(title="队列汇总")
        data = [[r.get(h) for h in QUEUE_HEADERS]
                for r in rows if isinstance(r, dict)]
        _write_table(ws, QUEUE_HEADERS, data, decimals=decimals, freeze=freeze,
                     autofilter=autofilter, width_max=width_max,
                     note="队列没有可汇总的目录")
        path = _save_workbook_atomic(wb, out_path)
        if path:
            log.info("已写出队列汇总: %s（%d 个目录）", path, len(data))
        return path
    except Exception as err:                            # noqa: BLE001 —— IO 层铁律：不抛
        log.warning("队列汇总导出失败（不影响队列结果本身）: %s", err)
        return None


def write_statistics_workbook(detail_rows, stat_rows, out_path,
                              cfg=None) -> Optional[Path]:
    """写出多 run 统计报表 xlsx（V2.5 FR-30.2：**明细 + 统计两 sheet**）。

    行输入来自 :func:`core.report.summary.collect_statistics` 的 detail/stat。
    空明细 → None（不产空文件，A30.3）；写盘失败 → None。
    """
    try:
        if not detail_rows:
            log.info("统计报表无明细行 → 不生成空文件")
            return None
        if not _ensure_openpyxl():
            log.warning("未安装 openpyxl，无法导出统计报表（pip install openpyxl）：%s",
                        _IMPORT_ERR)
            return None
        from .summary import STAT_DETAIL_HEADERS, STAT_HEADERS
        decimals, width_max, freeze, autofilter = _xlsx_opts(cfg or {})
        _styles()
        wb = Workbook()
        try:
            wb.remove(wb.active)
        except Exception:                               # noqa: BLE001
            pass
        ws_d = wb.create_sheet(title="明细")
        _write_table(ws_d, STAT_DETAIL_HEADERS,
                     [[r.get(h) for h in STAT_DETAIL_HEADERS]
                      for r in detail_rows if isinstance(r, dict)],
                     decimals=decimals, freeze=freeze, autofilter=autofilter,
                     width_max=width_max, note="没有可统计的 run")
        ws_s = wb.create_sheet(title="统计")
        _write_table(ws_s, STAT_HEADERS,
                     [[r.get(h) for h in STAT_HEADERS]
                      for r in (stat_rows or []) if isinstance(r, dict)],
                     decimals=decimals, freeze=freeze, autofilter=autofilter,
                     width_max=width_max, note="没有可聚合的组")
        path = _save_workbook_atomic(wb, out_path)
        if path:
            log.info("已写出统计报表: %s（明细 %d 行 / 统计 %d 行）",
                     path, len(detail_rows), len(stat_rows or []))
        return path
    except Exception as err:                            # noqa: BLE001
        log.warning("统计报表导出失败（不影响既有 run 产物）: %s", err)
        return None
