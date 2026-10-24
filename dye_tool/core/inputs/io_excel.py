# -*- coding: utf-8 -*-
"""Excel 工作簿（.xlsx）导入适配器（V2.6 FR-35 / 路线 C1）。

为什么单独一个模块
------------------
与 core.inputs.io_fds / core.inputs.io_spc 同理：io.py 承载文本格式解析与扫描
缓存，.xlsx 是 ZIP 容器（没有"行文本"概念）。保持"一个格式一个适配器"，由
io.py 的内容嗅探按需调用；默认关（io.excel.enable 为 False）时既有格式解析
路径**逐位不变**。

识别判据（FR-35.1）
-------------------
.xlsx 后缀 + ZIP 魔数 PK\x03\x04 + 容器内存在 xl/workbook.xml。
只看后缀不够（改扩展名的坏文件/加密存档都会被误认）；只看魔数也不够（任意
ZIP 都能过第一层）。三层同时满足才算 Excel 工作簿，坏文件一律返回中文原因，
**不抛异常**（由调用方转成文件表的"剔除原因"）。

解析语义（FR-35.4，写死规则 + WARNING）
----------------------------------------
* 取**首个**工作表（多 sheet 只记 WARNING，不猜数据在哪一页）；
* read_only=True, values_only=True, data_only=True（不执行公式，读缓存值）；
* 首行**没有任何数值** → 视为表头跳过；表头名匹配优先：
  - x 列：含 波长 / wavelength / lambda / λ，或整格等于 nm；
  - y 列：含 吸光度 / absorbance / intensity / 强度，或整格等于 abs / a；
  - 匹配不到 → 取首个数据行的**前两个数值列**；
* 坏行跳过（x 非数值）、空 y → NaN、±inf y → NaN；有效点 < 10 → 剔除；
* 数据行上限 200000 → 截断 + WARNING；
* 波长非严格升序 → 按波长稳定排序（与文本路径同策略）。

依赖：openpyxl **函数内懒导入**（先例 core/report/excel_export.py:71）；
缺失时给可操作的中文提示，绝不让整个程序 import 期就崩。不 import GUI /
matplotlib；不新增配置键（开关在 core.config.DEFAULT_CFG["io"]["excel"]）。
"""
from __future__ import annotations

import logging
import math
from pathlib import Path
from typing import Optional

import numpy as np

log = logging.getLogger("dye_tool.io.excel")

#: 数据点下限（与文本 / FDS / SPC 路径一致：有效点不足 10 视为不可用）
MIN_POINTS = 10

#: 数据行上限（FR-35.4：超过则截断 + WARNING，防止误选超大工作簿卡死）
MAX_ROWS = 200_000

#: 扩展名（内容嗅探只对 .xlsx 生效；.xls 老格式不在本批范围）
XLSX_SUFFIX = ".xlsx"

#: ZIP 本地文件头魔术字
ZIP_MAGIC = b"PK\x03\x04"

#: 工作簿部件名（ZIP 内存在它才说明这是 Excel 工作簿，而不是任意 ZIP）
_WORKBOOK_PART = "xl/workbook.xml"

#: 表头识别词（FR-35.4；统一小写匹配）
_X_CONTAINS = ("波长", "wavelength", "lambda", "λ")
_X_EXACT = {"nm"}
_Y_CONTAINS = ("吸光度", "absorbance", "intensity", "强度")
_Y_EXACT = {"abs", "a"}


class ExcelError(ValueError):
    """xlsx 解析失败。消息为**中文人话**，可直接作为文件表"剔除原因"。"""


# ═══════════════════════════ 内容嗅探 ═══════════════════════════
def looks_like_xlsx(data: bytes) -> bool:
    """字节流是否是 Excel 工作簿（ZIP 魔数 + xl/workbook.xml 存在）。"""
    if len(data) < 4 or not data.startswith(ZIP_MAGIC):
        return False
    import io as _io
    import zipfile

    try:
        with zipfile.ZipFile(_io.BytesIO(data)) as zf:
            names = set(zf.namelist())
    except Exception:                                   # noqa: BLE001 —— 坏 ZIP 不是 Excel
        return False
    return _WORKBOOK_PART in names


def is_excel_file(path: Path | str) -> bool:
    """按**后缀 + 内容**判断是不是 Excel 工作簿（读盘失败 / 坏文件 → False，不抛）。"""
    p = Path(str(path))
    if p.suffix.lower() != XLSX_SUFFIX:
        return False
    try:
        data = p.read_bytes()
    except OSError:
        return False
    return looks_like_xlsx(data)


def _import_openpyxl():
    """延迟 import openpyxl；缺失时给**可操作**的中文提示。"""
    try:
        import openpyxl
    except ImportError as err:                          # pragma: no cover
        raise ExcelError(
            "缺少 openpyxl 依赖，无法读取 .xlsx 文件"
            "（安装：pip install openpyxl）") from err
    return openpyxl


# ═══════════════════════════ 单元格 / 表头 ═══════════════════════════
def _num(value) -> Optional[float]:
    """单元格 → 有限 float；空 / 布尔 / 非数值 / inf → None。"""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        v = float(value)
        return v if math.isfinite(v) else None
    txt = str(value).strip().strip('"').strip("'")
    if not txt:
        return None
    try:
        v = float(txt)
    except ValueError:
        return None
    return v if math.isfinite(v) else None


def _find_header_col(header, contains, exact) -> Optional[int]:
    """在表头行里找第一个命中词表的列号（不含 → None）。"""
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        txt = str(cell).strip().lower()
        if not txt:
            continue
        if any(tok in txt for tok in contains) or txt in exact:
            return idx
    return None


def _numeric_cols(row) -> list:
    """一行里数值单元格的列号（升序）。"""
    return [idx for idx, value in enumerate(row) if _num(value) is not None]


def _normalize_ascending(wl: np.ndarray, a: np.ndarray) -> tuple:
    """波长非严格升序时按波长稳定排序（与文本路径同策略）。"""
    if wl.size < 2:
        return wl, a
    if np.all(np.diff(wl) > 0):
        return wl, a
    log.warning("Excel 波长列非严格升序 → 按波长升序重排")
    order = np.argsort(wl, kind="stable")
    return wl[order], a[order]


# ═══════════════════════════ 首 sheet 解析 ═══════════════════════════
def _read_sheet(ws) -> tuple[Optional[np.ndarray], Optional[np.ndarray], str]:
    """解析一个工作表 → (wl, a, 中文失败原因)；失败时前两项为 None。"""
    it = iter(ws.iter_rows(values_only=True))
    first = next(it, None)
    if first is None:
        return None, None, "首个工作表为空"
    header = None
    if not any(_num(value) is not None for value in first):
        header = first                              # 首行无任何数值 → 表头
    pending: list = [] if header is not None else [first]
    x_idx = y_idx = None
    if header is not None:
        x_idx = _find_header_col(header, _X_CONTAINS, _X_EXACT)
        y_idx = _find_header_col(header, _Y_CONTAINS, _Y_EXACT)
    # 列识别：表头匹配不到时，从数据行里找出前两个数值列。不能只看第一行——
    # 首行 y 可能为空（真实数据的 t=0 空白段）或整行只有一列；向后扫描到出现
    # "至少两列数值"的一行为止，扫过的行进入 pending，随后照常按数据行处理。
    if x_idx is None or y_idx is None:
        cols = None
        while True:
            row = next(it, None)
            if row is None:
                break
            pending.append(row)
            found = _numeric_cols(row)
            if len(found) >= 2:
                cols = found
                break
            if len(pending) >= MAX_ROWS:
                break
        if cols is None:
            return None, None, "首个工作表没有至少两列数值数据"
        if x_idx is None and y_idx is None:
            x_idx, y_idx = cols[0], cols[1]
        elif x_idx is None:
            x_idx = next((c for c in cols if c != y_idx), None)
        else:
            y_idx = next((c for c in cols if c != x_idx), None)
        if x_idx is None or y_idx is None:
            return None, None, "首个工作表没有至少两列数值数据"
    xs: list = []
    ys: list = []
    bad_x = skipped = inf_y = 0
    truncated = False

    def take(row) -> bool:
        """收一行；返回 False 表示已达行数上限（截断）。"""
        nonlocal bad_x, skipped, inf_y
        if len(xs) >= MAX_ROWS:
            return False
        x = _num(row[x_idx]) if x_idx < len(row) else None
        if x is None:
            bad_x += 1
            return True
        if y_idx >= len(row):
            # 行不完整（缺 y 列）：按坏行跳过；单元格存在但为空才是 NaN。
            bad_x += 1
            return True
        y_raw = row[y_idx]
        if isinstance(y_raw, float) and not math.isfinite(y_raw):
            inf_y += 1
        y = _num(y_raw)
        if y is None:
            skipped += 1
            y = float("nan")
        xs.append(x)
        ys.append(y)
        return True

    for row in pending:
        if not take(row):
            truncated = True
            break
    if not truncated:
        for row in it:
            if not take(row):
                truncated = True
                break
    if truncated:
        log.warning("Excel 数据行超过上限 %d，已截断（仅取前 %d 行）",
                    MAX_ROWS, MAX_ROWS)
    if bad_x:
        log.info("Excel 有 %d 行波长非数值已跳过", bad_x)
    if inf_y:
        log.warning("Excel 有 %d 个非有限吸光度（±inf）已归一为 NaN", inf_y)
    n = len(xs)
    if n < MIN_POINTS:
        return None, None, f"有效数据点不足({n} 点)"
    wl = np.asarray(xs, float)
    a = np.asarray(ys, float)
    wl, a = _normalize_ascending(wl, a)
    return wl, a, ""


# ═══════════════════════════ 主解析 ═══════════════════════════
def read_excel_file_ex(path: Path | str) -> tuple[Optional[list[dict]], str]:
    """读一个 .xlsx 文件 → (记录列表 | None, 中文失败原因)。

    成功返回一文件一条记录（一文件一 Scan）::

        [{"wl": ndarray, "a": ndarray, "source": "xlsx", "meta": {}}]

    失败（后缀不符 / 坏 ZIP / 无法打开 / 无数据 / 点数不足等）返回
    (None, 中文原因)；**不抛异常**（openpyxl 缺失也走该通道）。
    """
    p = Path(str(path))
    if p.suffix.lower() != XLSX_SUFFIX:
        return None, f"不是 .xlsx 文件：{p.name}"
    try:
        data = p.read_bytes()
    except OSError as err:
        return None, f"读取失败：{err}"
    if not looks_like_xlsx(data):
        return None, (f"不是有效的 .xlsx 文件：{p.name}"
                      "（ZIP 魔数或 xl/workbook.xml 缺失）")
    try:
        openpyxl = _import_openpyxl()
    except ExcelError as err:
        return None, str(err)
    wb = None
    try:
        try:
            wb = openpyxl.load_workbook(filename=str(p), read_only=True,
                                        data_only=True, keep_links=False)
        except Exception as err:                        # noqa: BLE001
            return None, (f"Excel 文件无法打开（可能加密或损坏）："
                          f"{type(err).__name__}: {err}")
        names = list(getattr(wb, "sheetnames", None) or [])
        if not names:
            return None, "工作簿没有工作表"
        if len(names) > 1:
            log.warning("%s 含 %d 个工作表，按 FR-35.4 只读取首个工作表 %r",
                        p.name, len(names), names[0])
        wl, a, why = _read_sheet(wb[names[0]])
        if wl is None:
            return None, why
        return [{"wl": wl, "a": a, "source": "xlsx", "meta": {}}], ""
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:                           # noqa: BLE001
                pass


def read_excel_file(path: Path | str):
    """读一个 .xlsx 文件 → (wl, a, meta)；**任何失败返回 None**。

    与 io_fds.read_fds_file / io_spc.read_spc_file 同构，便于在内容嗅探链里与
    既有格式并列；需要失败原因时用 read_excel_file_ex。
    """
    recs, _why = read_excel_file_ex(path)
    if not recs:
        return None
    rec = recs[0]
    return rec["wl"], rec["a"], dict(rec.get("meta") or {})


__all__ = [
    "ExcelError", "XLSX_SUFFIX", "ZIP_MAGIC", "MIN_POINTS", "MAX_ROWS",
    "looks_like_xlsx", "is_excel_file", "read_excel_file_ex", "read_excel_file",
]
