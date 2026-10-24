# -*- coding: utf-8 -*-
"""P5 Excel 多 sheet 导出（``core/excel_export.py``）的契约测试。

覆盖设计文档要求的 5 条：

1. 工作簿可被 openpyxl 重新读回，表名集合与 ``cfg["excel"]["sheets"]`` **精确对应**；
2. ``summary`` 里 k 的数值与 ``run_result`` 里的 k **一致到 ``decimals`` 位**
   （防"导出的数和算出的数不一样"——本文件最重要的一条）；
3. 关掉某张表后该表**确实不存在**；
4. 缺数据的假 run_result **不抛异常**且表里有"无数据"提示；
5. 列宽不超过 ``column_width_max``；冻结首行 / 自动筛选已生效。

数据用 :func:`core.synthetic.make_synthetic` 造的**真值已知**的合成数据，跑一次
**真实 pipeline**（不是手工塞假 results），因此导出链路与生产完全同源。
"""
from __future__ import annotations

import copy
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.config import default_config
from core.excel_export import SHEET_ORDER, export_workbook
from core.pipeline import RunResult, run_phi_pipeline
from core.synthetic import make_synthetic

NOTE = "（无数据"


# ═══════════════════════════ 夹具 ═══════════════════════════
@pytest.fixture(scope="module")
def synth(tmp_path_factory):
    """跑一次真实 pipeline（synthetic 数据）并导出工作簿；返回 ``(rr, cfg, path)``。

    模块级：pipeline 只跑一次，多条断言共用（pytest 默认 ``-m "not slow"`` 也会
    跑它，故用 ``only={"a"}`` + 预览档把出图压到最小）。
    """
    base = tmp_path_factory.mktemp("p5_excel")
    data = base / "data"
    make_synthetic(data, seed=42)
    cfg = default_config()
    # 与 core.synthetic.self_test 同一套真值设定：synthetic 的 λ_ex 处吸光度靠
    # a_ex_auto 取不准（基线校正后为负），显式声明后 Φ_Δ 才算得出来。
    cfg["quantum_yield"]["a_ex"] = {"sample": 0.15, "ref": 0.15}
    cfg["quantum_yield"]["phi_ref"] = 0.75
    rr = run_phi_pipeline(data, cfg, out_root=base / "runs",
                          export_mode="preview", only={"a"})
    assert rr.ok is True and rr.run_dir is not None
    path = export_workbook(rr, cfg)
    assert path is not None and Path(path).exists()
    return rr, cfg, Path(path)


def _head(ws) -> list:
    return [c.value for c in ws[1]]


def _rows(ws) -> list:
    """数据行（含"无数据"提示行），每行 = ``dict(列名 → 值)``。"""
    head = _head(ws)
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def _has_note(ws) -> bool:
    return any(NOTE in str(c.value) for row in ws.iter_rows() for c in row
               if c.value is not None)


# ═══════════════════════════ 1. 表名与配置精确对应 ═══════════════════════════
def test_sheet_names_match_config_exactly(synth) -> None:
    """回读工作簿：表名集合 == ``cfg["excel"]["sheets"]`` 的键集合。"""
    _rr, cfg, path = synth
    wb = load_workbook(path)
    assert set(wb.sheetnames) == set(cfg["excel"]["sheets"]) == set(SHEET_ORDER)
    # 每张表都有表头（列契约不因数据多少而变），且表头非空、无重复
    for name in wb.sheetnames:
        head = _head(wb[name])
        assert head and all(h is not None and str(h).strip() for h in head), name
        assert len(set(head)) == len(head), f"{name} 表头有重复列名"


def test_required_columns_are_present(synth) -> None:
    """设计文档点名的列必须都在。"""
    _rr, _cfg, path = synth
    wb = load_workbook(path)
    for col in ("染料", "组", "时间点数", "k (s⁻¹)", "k 标准误", "R²", "Φ_Δ",
                "Φ_Δ 误差", "参比", "状态", "备注", "run 目录"):
        assert col in set(_head(wb["summary"])), col
    for col in ("文件", "组", "t", "A0", "A0 池化", "是否饱和", "异常点数",
                "无效点数", "剔除原因", "QC 状态"):
        assert col in set(_head(wb["qc"])), col
    for col in ("t", "A/A0", "是否参与拟合", "是否被截断"):
        assert col in set(_head(wb["kinetics"])), col
    for col in ("k (s⁻¹)", "k 标准误", "R²", "Φ_Δ", "Φ_Δ 误差", "run 时间"):
        assert col in set(_head(wb["dyes"])), col
    assert _head(wb["params"]) == ["参数", "值"]


# ═══════════════════════════ 2. summary 的数 == run_result 的数 ═══════════════════════════
def test_summary_k_matches_run_result_to_decimals(synth) -> None:
    """**核心断言**：导出的 k / k 标准误 / R² / n / Φ 与 run_result 一致到 ``decimals`` 位。"""
    rr, cfg, path = synth
    d = int(cfg["excel"]["decimals"])
    tol = 0.5 * 10 ** (-d)
    ws = load_workbook(path)["summary"]
    table = {r["组"]: r for r in _rows(ws)}
    assert set(table) == set(rr.results)          # 每组一行，不多不少

    n_k = 0
    for g, e in rr.results.items():
        row = table[g]
        f = e["pooled"]
        # 时间点数 == reps 里实际数组长度之和（不是编出来的）
        assert row["时间点数"] == sum(int(np.asarray(r["t"]).size)
                                      for r in e["reps"].values())
        if f is None or f.status == "FAILED" or not np.isfinite(f.k):
            assert row["k (s⁻¹)"] is None          # 拟合失败 → 空格，不写假数
            continue
        cell = row["k (s⁻¹)"]
        assert isinstance(cell, float) and not isinstance(cell, bool)
        assert abs(cell - f.k) <= tol, (g, cell, f.k)
        assert round(cell, d) == round(f.k, d)
        assert abs(row["k 标准误"] - f.k_err) <= tol
        assert abs(row["R²"] - f.r2) <= tol
        assert row["n"] == f.n
        assert row["截断"] == ("是" if f.truncated else "否")
        assert row["状态"] == f.status
        assert row["角色"] == e["role"]
        n_k += 1
    assert n_k >= 5                                # synthetic 里至少 6 组可拟合

    # Φ_Δ 列 == compute_phi 的结果（含"无参比/失败"组）
    assert rr.phi, "本夹具应至少算出一组 Φ"
    for g, p in rr.phi.items():
        row = table[g]
        if p.get("status") == "OK":
            assert abs(row["Φ_Δ"] - p["phi"]) <= tol
            assert abs(row["Φ_Δ 误差"] - p["phi_err"]) <= tol
            assert row["参比"] == "、".join(q["ref"] for q in p["pairs"])
        else:
            assert row["Φ_Δ"] is None
            assert row["参比"] == p["status"]
    assert "bg" not in table                       # 背景组不进结果表


def test_kinetics_flags_match_fit_inputs(synth) -> None:
    """kinetics 明细：参与拟合的点数 == 该重复拟合的 n；截断组确实标了"被截断"。"""
    rr, _cfg, path = synth
    rows = _rows(load_workbook(path)["kinetics"])
    per: dict = {}
    for r in rows:
        per.setdefault((r["组"], r["重复"]), []).append(r)
    for g, e in rr.results.items():
        for rep, r in e["reps"].items():
            key = (g, rep)
            assert key in per, key
            fit = r["fit"]
            used = sum(1 for x in per[key] if x["是否参与拟合"] == "是")
            if fit is not None and fit.status != "FAILED":
                assert used == fit.n, key
            else:
                assert used == 0, key
    # 被截断舍弃的点：必须"不参与拟合"且"是否被截断=是"
    cut = [r for r in rows if r["是否被截断"] == "是"]
    assert cut and all(r["是否参与拟合"] == "否" for r in cut)
    assert {r["组"] for r in cut} == {g for g, e in rr.results.items()
                                      if e["pooled"] is not None
                                      and e["pooled"].truncated}
    # 数值列是数值（不是字符串），参与拟合的点 R 必须有限
    assert all(isinstance(r["t"], (int, float)) for r in rows)
    assert all(np.isfinite(r["A/A0"]) for r in rows if r["是否参与拟合"] == "是")


# ═══════════════════════════ 3. 关掉的表不进工作簿 ═══════════════════════════
def test_disabled_sheets_are_absent(synth, tmp_path) -> None:
    """把 kinetics / dyes 关掉 → 工作簿里确实没有这两张表，其余照常。"""
    rr, cfg, _path = synth
    cfg2 = copy.deepcopy(cfg)
    cfg2["excel"]["sheets"]["kinetics"] = False
    cfg2["excel"]["sheets"]["dyes"] = False
    p = export_workbook(rr, cfg2, out_path=tmp_path / "some_off.xlsx")
    assert p is not None and Path(p).exists()
    wb = load_workbook(p)
    assert set(wb.sheetnames) == {"summary", "qc", "params"}
    assert wb["summary"].max_row > 1               # 其余表仍有数据
    # 只留 summary 也一样
    cfg3 = copy.deepcopy(cfg)
    cfg3["excel"]["sheets"] = {k: (k == "summary") for k in SHEET_ORDER}
    p3 = export_workbook(rr, cfg3, out_path=tmp_path / "only_summary.xlsx")
    assert load_workbook(p3).sheetnames == ["summary"]


def test_disabled_by_enabled_flag(synth, tmp_path) -> None:
    """``excel.enabled=false`` → 不写文件、返回 None、不抛异常。"""
    rr, cfg, _path = synth
    cfg2 = copy.deepcopy(cfg)
    cfg2["excel"]["enabled"] = False
    out = tmp_path / "should_not_exist.xlsx"
    assert export_workbook(rr, cfg2, out_path=out) is None
    assert not out.exists()


def test_all_sheets_off_yields_no_broken_workbook(synth, tmp_path) -> None:
    """所有表都关掉 → 返回 None（不生成 Excel 打不开的空工作簿）。"""
    rr, cfg, _path = synth
    cfg2 = copy.deepcopy(cfg)
    cfg2["excel"]["sheets"] = {k: False for k in SHEET_ORDER}
    out = tmp_path / "empty_wb.xlsx"
    assert export_workbook(rr, cfg2, out_path=out) is None
    assert not out.exists()


# ═══════════════════════════ 4. 缺数据不抛异常 + 提示行 ═══════════════════════════
def test_missing_data_writes_note_and_never_raises(tmp_path) -> None:
    """空 run_result：不抛异常，每张（本该有数据的）表都写"（无数据：…）"。"""
    cfg = default_config()
    cfg["gui"]["out_root"] = str(tmp_path / "no_such_runs")     # 别去扫真 runs/
    out = tmp_path / "empty_run.xlsx"
    rr = RunResult()                                   # results/scans/run_dir 全空
    p = export_workbook(rr, cfg, out_path=out)
    assert p is not None and Path(p).exists()
    wb = load_workbook(p)
    assert set(wb.sheetnames) == set(SHEET_ORDER)
    for name in ("summary", "qc", "kinetics", "dyes"):
        ws = wb[name]
        assert _has_note(ws), f"{name} 缺数据时应写提示行"
        assert str(ws.cell(row=2, column=1).value).startswith(NOTE)
        assert _head(ws)                               # 表头仍在（列契约不失效）
        assert ws.max_row == 2                         # 只有提示行，没有假数据


def test_hostile_run_result_never_raises(tmp_path) -> None:
    """残缺对象 / None / 字段类型错乱 / 写盘失败 → 一律返回 None 或空表，绝不抛。"""
    cfg = default_config()
    cfg["gui"]["out_root"] = str(tmp_path / "no_such_runs")
    bad_inputs = [None, object(), {"results": None, "scans": "not-a-list"},
                  RunResult(results=None, scans=None, phi=None, controls=None)]
    for i, bad in enumerate(bad_inputs):
        out = tmp_path / f"bad_{i}.xlsx"
        p = export_workbook(bad, cfg, out_path=out)
        assert p is None or Path(p).exists()
    # 写盘失败（out_path 是**目录**）→ 顶层兜底返回 None
    assert export_workbook(RunResult(), cfg, out_path=tmp_path) is None
    # run_dir 与 out_path 都缺 → 无法确定输出位置 → None（不抛）
    assert export_workbook(RunResult(), cfg) is None
    # cfg 残缺（连 excel 节都没有）→ **仍按"全部生成"出一本**（L20 修复）
    # 原实现把"缺 sheets 键"当成"所有表都关掉"，返回 None 并打一条误导日志
    # （说成是用户关掉了），与 `enabled` 默认 True 的取向不对称。现在两端一致。
    p_nocfg = export_workbook(RunResult(), {}, out_path=tmp_path / "nocfg.xlsx")
    assert p_nocfg is not None and Path(p_nocfg).exists()
    assert set(load_workbook(p_nocfg).sheetnames) == set(SHEET_ORDER)
    # 但**显式全关**仍然返回 None（那是用户的明确意图）
    off = default_config()
    off["excel"]["sheets"] = {k: False for k in SHEET_ORDER}
    assert export_workbook(RunResult(), off, out_path=tmp_path / "off.xlsx") is None
    # 字段类型错乱但不为空 → 仍能出一本（内容为空表），不抛
    weird = RunResult(results=["nope"], scans=123, phi="x", controls=7)
    p = export_workbook(weird, cfg, out_path=tmp_path / "weird.xlsx")
    assert p is not None and Path(p).exists()
    for name in ("summary", "qc", "kinetics"):
        assert _has_note(load_workbook(p)[name]), name


# ═══════════════ L19/G10：超上限截断（第二轮审查 G10+G12） ═══════════════
def test_write_table_truncates_within_excel_row_limit(monkeypatch) -> None:
    """L19 的约定测试（G12 指出原先对 ``XLSX_MAX_ROWS`` **零覆盖**）。

    **G10（第二轮审查，修复轮新引入）**：原先 `limit = XLSX_MAX_ROWS - 1` **差一行**
    ——表头 1 行 + limit 行 = 恰好顶格 1,048,576；再追加"已截断"说明行就是
    **第 1,048,577 行，超上限**，正是 L19 要防的形态（Excel 弹"是否修复"；
    若 openpyxl 改为抛错，本函数在每表 try 之外，整本导出会中止）。
    现在 `- 2`：表头 + 数据 + 说明行 = 1,048,576，刚好顶格不超。

    上限被 patch 成 10 行，**跑起来很快**，而断言的正是同一套算术。
    """
    from openpyxl import Workbook

    from core import excel_export as ex

    monkeypatch.setattr(ex, "XLSX_MAX_ROWS", 10)        # 上限 10：表头+数据+说明
    ws = Workbook().active
    n = ex._write_table(ws, ["a", "b"],
                        [[float(i), f"t{i}"] for i in range(50)],
                        decimals=6, freeze=True, autofilter=True, width_max=60)
    # 数据行 = 上限 - 2（给表头与说明行各留 1 行）
    assert n == 8, f"应写 8 行数据（10-2），实际 {n}"
    assert ws.max_row == 10, f"总行数必须恰好顶格 10，实际 {ws.max_row}"
    assert ws.max_row <= 10, "**超出 Excel 上限** —— 这正是 G10 的形态"
    last = str(ws.cell(row=ws.max_row, column=1).value)
    assert "截断" in last or "保留" in last, f"说明行应在最后一行，实际 {last!r}"

    # 不超限时不截断、也不该有说明行
    ws2 = Workbook().active
    n2 = ex._write_table(ws2, ["a"], [[1.0]] * 3, decimals=6, freeze=True,
                         autofilter=True, width_max=60)
    assert n2 == 3 and ws2.max_row == 4, "未超限不该截断"
    assert all("截断" not in str(c.value) for row in ws2.iter_rows() for c in row)


# ═══════════════════════════ 5. 版式：列宽封顶 / 冻结 / 筛选 ═══════════════════════════
def test_width_cap_freeze_and_autofilter(synth) -> None:
    """每张表：冻结 A2、筛选覆盖整表、列宽不超过上限。"""
    _rr, cfg, path = synth
    wb = load_workbook(path)
    cap = float(cfg["excel"]["column_width_max"])
    widths: list = []
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.freeze_panes == "A2", name
        assert ws.auto_filter.ref == (
            f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"), name
        widths += [d.width for d in ws.column_dimensions.values()
                   if d.width is not None]
    assert widths and max(widths) <= cap + 1e-9


def test_width_cap_is_really_enforced(synth, tmp_path) -> None:
    """把上限压到 12：所有列宽 ≤ 12，且**至少有一列真的被截到 12**
    （表头「辐照度 (mW/cm²)」显示宽度 17 > 12，必然触发封顶）。"""
    rr, cfg, _path = synth
    cfg2 = copy.deepcopy(cfg)
    cfg2["excel"]["column_width_max"] = 12
    p = export_workbook(rr, cfg2, out_path=tmp_path / "narrow.xlsx")
    wb = load_workbook(p)
    widths = [d.width for ws in wb.worksheets
              for d in ws.column_dimensions.values() if d.width is not None]
    assert widths and all(w <= 12 + 1e-9 for w in widths)
    assert max(widths) == pytest.approx(12.0)


def test_layout_switches_can_be_turned_off(synth, tmp_path) -> None:
    """``freeze_header`` / ``autofilter`` 关掉后确实不生效。"""
    rr, cfg, _path = synth
    cfg2 = copy.deepcopy(cfg)
    cfg2["excel"]["freeze_header"] = False
    cfg2["excel"]["autofilter"] = False
    p = export_workbook(rr, cfg2, out_path=tmp_path / "nofreeze.xlsx")
    ws = load_workbook(p)["summary"]
    assert ws.freeze_panes is None
    assert ws.auto_filter.ref is None


# ═══════════════════════════ 中文 / 数值落盘确认 ═══════════════════════════
def test_chinese_headers_roundtrip_as_unicode(synth) -> None:
    """中文表头与说明回读后仍是原字符（UTF-8 写入，不该被转义）。"""
    _rr, _cfg, path = synth
    wb = load_workbook(path)
    head = _head(wb["summary"])
    assert "k (s⁻¹)" in head and "Φ_Δ 误差" in head and "辐照度 (mW/cm²)" in head
    notes = [str(r["说明"]) for r in _rows(wb["kinetics"])]
    assert any("超出拟合窗" in n for n in notes)       # 中文 + 数值混排
    reasons = [str(r["剔除原因"]) for r in _rows(wb["qc"])]
    assert any("波长轴" in x for x in reasons)         # synthetic 的 mis_1 整组被剔


def test_qc_sheet_flags_the_injected_spike(synth) -> None:
    """synthetic 在 sample_1 的 t=150 s、379 nm 处注入了尖峰（正好在探针窗口内）：
    QC 表应把它计成该文件的 1 个异常点，并给出备注；该文件 t 也对得上。"""
    _rr, cfg, path = synth
    assert float(cfg["wavelength"]["probe"]) == 379.0      # 尖峰就注入在探针波长上
    qc = {r["文件"]: r for r in _rows(load_workbook(path)["qc"])}
    key = "sample_1-150.txt"
    assert key in qc, sorted(qc)[:5]
    row = qc[key]
    assert row["t"] == 150
    assert row["组"] == "sample"
    assert row["异常点数"] == 1
    assert row["无效点数"] >= 1
    assert "探针窗口内 1 个无效点" in str(row["备注"])
    # 正常文件不该被误报
    clean = qc["sample_1-0.txt"]
    assert clean["异常点数"] == 0 and clean["QC 状态"] != "剔除"
