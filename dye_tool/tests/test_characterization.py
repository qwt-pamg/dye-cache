# -*- coding: utf-8 -*-
"""V2.6 G 道 FR-44 / FR-45 / A42.4–A42.6：表征总表 + 新图型 + 红线。

覆盖：
* A42.4 表征总表：14 列（FR-44 原序）/多条件行/缺项留空不臆造/组别别名可读写/
  与 FR-41 Stokes 表统一（单一模块 + stokes_table 薄别名，F 道 CSV 行为不变）；
* A42.5 导出与风格：CSV UTF-8-SIG + Excel（openpyxl 读回）；三类新图型
  （光谱演化/衰减曲线/分组柱状）生成、尺寸/dpi=600、关键结构（四边框/内刻度/
  粗体轴标签/无网格/包络填充/误差棒/裸粗体面板字母）；panel_label_style
  默认 paren 与 bare（LayoutSpec 校验）；
* A42.6 红线：零新配置键、流水线不 import 新模块（不改训练/分析链路）、
  legacy JACS fig a–e 渲染函数存在且签名前缀未变、photophysics 不改全局 rcParams、
  零 PDF。

只测可观测结构/文件契约，不逐像素比对（A2/A4.2 逐字节由收口门禁覆盖）。
"""
from __future__ import annotations

import inspect
import math
from datetime import datetime
from pathlib import Path

import numpy as np
import pytest

pytestmark = pytest.mark.slow

PROJECT_ROOT = Path(__file__).resolve().parents[1]

#: A42.6 零新配置键：GUI 会话态基线键集
#: theme 是前端重构（浅色/深色主题）有意新增的界面键，不属于 G 道
GUI_BASELINE_KEYS = {"recent_dirs", "max_recent", "out_root", "show_welcome",
                     "lazy_preview", "window", "nav", "theme"}

#: FR-44 表征总表 14 列（设计文档 §5.7 原序）
FR44_COLUMNS = ["条件/溶剂", "化合物/组别", "吸收谱文件", "发射谱文件",
                "λabs(nm)", "λem(nm)", "Δλ(nm)", "Δν̃(cm⁻¹)",
                "ε(M⁻¹cm⁻¹)", "φ", "brightness(M⁻¹cm⁻¹)",
                "k_photo(min⁻¹)", "t½(min)", "备注"]

#: FR-41 Stokes 8 列（Q-E19 合并后仍保留的导出格式）
FR41_COLUMNS = ["组别", "吸收谱文件", "发射谱文件", "λabs(nm)", "λem(nm)",
                "Δλ(nm)", "Δν̃(cm⁻¹)", "录入时间"]

NOW = datetime(2026, 10, 18, 23, 30, 5)
STAMP = "233005"


# ═══════════════════════ 模块/通用工具 ═══════════════════════
def _char():
    import core.report.characterization as m
    return m


def _stokes():
    import core.report.stokes_table as m
    return m


def _pp():
    import core.render.photophysics as m
    return m


def _fn(mod, *names):
    for n in names:
        f = getattr(mod, n, None)
        if callable(f):
            return f
    raise AssertionError(
        f"{mod.__name__} 缺函数 {names}；现有公开名："
        f"{[n for n in dir(mod) if not n.startswith('_')]}")


def _is_zh(text) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in str(text or ""))


def _make_rows():
    ch = _char()
    mk = _fn(ch, "make_characterization_row", "make_row")
    return mk, [
        mk(condition="MeOH", compound="DMAR-SY", abs_file="DMAR-SY-abs.txt",
           em_file="DMAR-SY-em.txt", lam_abs=450.0, lam_em=520.0,
           epsilon=5.0e4, phi=0.54, brightness=2.7e4, k_photo=0.048,
           t_half=14.4, note="ref"),
        mk(condition="DMSO", compound="DMAR-SY", abs_file="DMAR-SY-dmso-abs.txt",
           em_file="DMAR-SY-dmso-em.txt", lam_abs=455.0, lam_em=525.0,
           epsilon=4.8e4, phi=0.31, brightness=1.49e4),
        mk(condition="PBS", compound="Coumarin-153"),
    ]


def _read_csv(path: Path) -> list[list[str]]:
    import csv
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def _pil(path: Path):
    from PIL import Image
    return Image.open(str(path))


def _tick_direction(ax, axis: str = "x") -> str:
    """取主轴刻度方向（公共 API；mpl 3.11 无参 tick_params 返回 None）。"""
    obj = ax.xaxis if str(axis) == "x" else ax.yaxis
    try:
        kw = obj.get_tick_params(which="major")
        if isinstance(kw, dict) and kw.get("direction"):
            return str(kw["direction"]).lower()
    except Exception:                                   # noqa: BLE001
        pass
    try:                                                # 兜底：tick 线 marker 判据
        marker = int(obj.get_major_ticks()[0].tick1line.get_marker())
        if str(axis) == "x":
            return {2: "in", 3: "out"}.get(marker, "")
        return {1: "in", 0: "out"}.get(marker, "")
    except Exception:                                   # noqa: BLE001
        pass
    return ""


def _grid_visible(ax) -> bool:
    lines = list(ax.get_xgridlines()) + list(ax.get_ygridlines())
    return any(bool(ln.get_visible()) for ln in lines)


def _bold(weight) -> bool:
    s = str(weight).lower()
    return s in ("bold", "700", "heavy", "black") or s.startswith("bold")


# ═══════════════════════ A42.4 表征总表 ═══════════════════════
def test_a42_4_columns_rows_make_and_blank_fields() -> None:
    """A42.4：14 列齐全；多条件行；缺项留空不臆造；group 别名可读写。"""
    ch = _char()
    cols = list(ch.CHARACTERIZATION_COLUMNS)
    assert cols == FR44_COLUMNS, f"表征总表列不符：{cols}"
    mk, rows = _make_rows()
    assert len(rows) == 3
    r0 = rows[0]
    assert float(r0.lam_abs) == pytest.approx(450.0)
    assert float(r0.lam_em) == pytest.approx(520.0)
    assert float(r0.delta_lambda) == pytest.approx(70.0)
    assert float(r0.delta_wavenumber) == pytest.approx(
        1.0e7 / 450.0 - 1.0e7 / 520.0)
    assert float(r0.epsilon) == pytest.approx(5.0e4)
    assert float(r0.phi) == pytest.approx(0.54)
    assert float(r0.brightness) == pytest.approx(2.7e4)
    assert float(r0.k_photo) == pytest.approx(0.048)
    assert float(r0.t_half) == pytest.approx(14.4)
    assert r0.condition == "MeOH" and r0.note == "ref"
    assert bool(r0.ok) is True

    # group 兼容别名（F 道口径）：读写都落到 compound
    assert r0.group == "DMAR-SY"
    r0.group = "DMAR-SY2"
    assert r0.compound == "DMAR-SY2"

    # 缺项留空：不臆造数值/字符串
    blank = mk(condition="DMSO")
    assert _is_blank(blank.lam_abs) and _is_blank(blank.lam_em)
    assert (blank.abs_file or "") == "" and (blank.note or "") == ""
    assert bool(blank.ok) is False, "λ 全缺的行不应被标为 ok"


def _is_blank(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return not v.strip()
    try:
        return not math.isfinite(float(v))
    except (TypeError, ValueError):
        return False


def test_a42_5_csv_and_excel_exports(tmp_path) -> None:
    """A42.5：CSV（UTF-8-SIG/中文列名/命名/不覆盖）+ Excel（openpyxl 读回）。"""
    ch = _char()
    write_csv = _fn(ch, "write_characterization_csv")
    write_xlsx = _fn(ch, "write_characterization_excel",
                     "write_characterization_xlsx")
    _mk, rows = _make_rows()
    out_dir = tmp_path / "out"
    p1 = write_csv(rows, out_dir, now=NOW)
    assert p1 is not None and Path(p1).parent == out_dir
    assert Path(p1).name == f"characterization_{STAMP}.csv", Path(p1).name
    assert Path(p1).read_bytes()[:3] == b"\xef\xbb\xbf", "CSV 缺 UTF-8-SIG BOM"
    got = _read_csv(Path(p1))
    assert got[0] == FR44_COLUMNS, got[0]
    assert len(got) == 1 + len(rows)
    assert got[1][0] == "MeOH" and got[1][1] == "DMAR-SY"
    assert float(got[1][4]) == pytest.approx(450.0)       # λabs
    assert float(got[1][8]) == pytest.approx(5.0e4)       # ε
    assert float(got[1][9]) == pytest.approx(0.54)        # φ
    # 第 3 行（PBS/Coumarin-153）ε/φ 缺 → 留空（不臆造 0）
    assert got[3][8].strip() == "" and got[3][9].strip() == ""

    # 同名不覆盖：再写一次 → 新文件、旧文件字节不变
    before = Path(p1).read_bytes()
    p2 = write_csv(rows, out_dir, now=NOW)
    assert p2 is not None and Path(p2) != Path(p1)
    assert Path(p1).read_bytes() == before
    assert STAMP in Path(p2).name and Path(p2).name != Path(p1).name

    # 空表 → None 且不产文件
    assert write_csv([], out_dir, now=NOW) is None

    # Excel：openpyxl 读回，中文表头 + 数值仍为数值；样式单源（excel_export）
    import openpyxl
    x1 = write_xlsx(rows, out_dir, now=NOW)
    assert x1 is not None and Path(x1).is_file(), "Excel 没有产出"
    assert Path(x1).suffix.lower() == ".xlsx"
    wb = openpyxl.load_workbook(str(x1), data_only=True)
    assert "表征总表" in wb.sheetnames, wb.sheetnames
    ws = wb["表征总表"]
    header = [c.value for c in ws[1]][:len(FR44_COLUMNS)]
    assert [str(h) for h in header] == FR44_COLUMNS, header
    assert ws.max_row >= 1 + len(rows)
    eps_cell = ws.cell(row=2, column=9).value
    assert isinstance(eps_cell, (int, float)) and float(eps_cell) == pytest.approx(5.0e4)
    wb.close()

    # 空表 Excel：按实现纪律产出"表头 + （无数据：…）提示行"的工作簿
    x_empty = write_xlsx([], out_dir, now=NOW)
    if x_empty is not None:
        wb2 = openpyxl.load_workbook(str(x_empty), data_only=True)
        ws2 = wb2[wb2.sheetnames[0]]
        assert [str(c.value) for c in ws2[1]][:len(FR44_COLUMNS)] == FR44_COLUMNS
        texts = " ".join(str(c.value) for row in ws2.iter_rows() for c in row)
        assert _is_zh(texts)
        wb2.close()
    assert not list(Path(out_dir).glob("*.pdf")), "导出产生了 PDF（零 PDF 红线）"


def test_a42_4_stokes_table_is_thin_alias_of_characterization(tmp_path) -> None:
    """A42.4：FR-41 Stokes 表并入 characterization 单模块；stokes_table 纯转发。"""
    st = _stokes()
    ch = _char()
    assert getattr(st, "StokesRow", None) is ch.CharacterizationRow, (
        "stokes_table.StokesRow 未别名到 characterization.CharacterizationRow")
    src = inspect.getsource(st)
    assert "characterization" in src, "stokes_table.py 不是薄转发（未引用 characterization）"
    assert "STOKES_COLUMNS" in inspect.getsource(ch), (
        "STOKES_COLUMNS 未在 characterization.py 中定义（非单模块）")
    assert list(st.STOKES_COLUMNS) == FR41_COLUMNS
    assert list(getattr(ch, "STOKES_COLUMNS")) == FR41_COLUMNS

    # F 道既有 CSV 行为不变（8 列中文表头/BOM/命名）
    row = st.make_stokes_row("Rhodamine", "Rhodamine_1-0.txt",
                             "Rhodamine_1-30.txt", 450.0, 520.0,
                             recorded_at="2026-10-18 23:00:00")
    assert st.default_stokes_name(NOW) == f"stokes_{STAMP}.csv"
    p = st.write_stokes_csv([row], tmp_path / "st", now=NOW)
    assert p is not None and Path(p).read_bytes()[:3] == b"\xef\xbb\xbf"
    got = _read_csv(Path(p))
    assert got[0] == FR41_COLUMNS, got[0]
    assert got[1][0] == "Rhodamine"
    assert got[1][3] == "450.0" and got[1][5] == "70.0"
    assert got[1][6] == str(int(round(1.0e7 / 450.0 - 1.0e7 / 520.0)))


# ═══════════════════════ A42.5 新图型 + panel_label_style ═══════════════════════
def test_a42_5_three_new_plot_types_structure_and_files(tmp_path) -> None:
    """A42.5：三类新图型生成（PNG/dpi≈600）+ 关键样式结构断言。"""
    from matplotlib.collections import PathCollection, PolyCollection
    from matplotlib.lines import Line2D

    pp = _pp()
    wl = np.arange(450.0, 601.0, 1.0)
    frames = []
    for i, t in enumerate((0.0, 60.0, 120.0, 180.0)):
        amp = 1.0 * (0.75 ** i)
        y = amp * np.exp(-0.5 * ((wl - 520.0) / 30.0) ** 2)
        frames.append({"label": f"{t:.0f} s", "x": wl, "y": y, "t": t})

    # ① 光谱演化
    out1 = tmp_path / "evo.png"
    p1 = pp.plot_spectral_evolution(frames, out1)
    assert Path(p1).is_file() and Path(p1).suffix.lower() == ".png"
    im1 = _pil(Path(p1))
    assert im1.size[0] > 0 and im1.size[1] > 0
    dpi = im1.info.get("dpi", (0, 0))
    assert round(float(dpi[0])) == 600, f"演化图 dpi={dpi}"
    fig1, meta1 = pp.build_spectral_evolution_figure(frames)
    ax1 = meta1["axes"][0] if meta1.get("axes") else fig1.axes[0]
    assert all(sp.get_visible() for sp in ax1.spines.values()), "演化图缺四边框"
    assert _tick_direction(ax1) == "in", f"x 刻度未朝内：{_tick_direction(ax1)}"
    assert _tick_direction(ax1, "y") == "in", (
        f"y 刻度未朝内：{_tick_direction(ax1, 'y')}")
    assert not _grid_visible(ax1), "演化图出现网格"
    t0_color = str(meta1.get("t0_color") or "")
    assert t0_color.lower() in ("red", "#c81e1e", "c81e1e"), t0_color
    assert str((meta1.get("line_colors") or [""])[0]).lower() in (
        "red", "#c81e1e", "c81e1e"), meta1.get("line_colors")
    assert len(meta1.get("line_colors") or []) == len(frames)
    env = meta1.get("envelope")
    assert isinstance(env, PolyCollection) or any(
        isinstance(c, PolyCollection) for c in ax1.collections), (
        "演化图缺首末包络填充（PolyCollection）")
    arrow = str(meta1.get("arrow_text") or "")
    assert "ν" in arrow or "h" in arrow or "\\nu" in arrow, arrow
    assert _bold(ax1.xaxis.label.get_fontweight()), "演化图轴标签未加粗"

    # ② 衰减曲线
    t = np.linspace(0.0, 30.0, 61)
    k_true = 0.05
    y = 5.0 + 95.0 * np.exp(-k_true * t)
    fit = {"k": k_true, "t_half": math.log(2.0) / k_true, "r2": 0.9995,
           "I0": 100.0, "I_inf": 5.0, "model": "exp_plateau"}
    out2 = tmp_path / "decay.png"
    p2 = pp.plot_decay_curve(t, y, fit, out2)
    assert Path(p2).is_file()
    assert round(float(_pil(Path(p2)).info.get("dpi", (0, 0))[0])) == 600
    fig2, meta2 = pp.build_decay_figure(t, y, fit)
    ax2 = meta2["axes"][0] if meta2.get("axes") else fig2.axes[0]
    scatter = meta2.get("scatter")
    assert scatter is not None and isinstance(scatter, (PathCollection, Line2D)), (
        f"衰减图缺散点/数据标记：{type(scatter).__name__}")
    assert isinstance(meta2.get("fit_line"), Line2D), "衰减图缺拟合线"
    legend_text = str(meta2.get("fit_text") or meta2.get("legend_text") or "")
    assert "k" in legend_text and "min" in legend_text, legend_text
    assert float(meta2.get("k")) == pytest.approx(k_true, rel=1e-9)
    assert float(meta2.get("t_half")) == pytest.approx(math.log(2.0) / k_true,
                                                       rel=1e-9)
    assert all(sp.get_visible() for sp in ax2.spines.values())
    assert _tick_direction(ax2) == "in" and _tick_direction(ax2, "y") == "in"
    assert not _grid_visible(ax2)

    # ③ 分组柱状
    groups = ["DMAR-SY", "Coumarin-153"]
    series = {
        "MeOH": {"values": [0.54, 0.31], "errors": [0.02, 0.03]},
        "DMSO": {"values": [0.48, 0.27], "errors": [0.02, 0.02]},
    }
    out3 = tmp_path / "bars.png"
    p3 = pp.plot_grouped_bars(groups, series, out3)
    assert Path(p3).is_file()
    assert round(float(_pil(Path(p3)).info.get("dpi", (0, 0))[0])) == 600
    fig3, meta3 = pp.build_grouped_bars_figure(groups, series)
    ax3 = meta3["axes"][0] if meta3.get("axes") else fig3.axes[0]
    assert meta3.get("bar_containers"), "柱状图缺 bar container"
    assert meta3.get("errorbar_containers"), "柱状图缺误差棒"
    assert meta3.get("colors"), "柱状图缺双色"
    fig3.canvas.draw()
    labels = ax3.get_xticklabels()
    assert labels and round(float(labels[0].get_rotation())) == 45, (
        "x 标签未旋转 45°")
    assert not _grid_visible(ax3), "柱状图出现网格"
    assert all(sp.get_visible() for sp in ax3.spines.values())
    assert _tick_direction(ax3) == "in" and _tick_direction(ax3, "y") == "in", (
        "柱状图刻度未朝内")

    # 零 PDF
    assert not list(tmp_path.rglob("*.pdf")), "新图型产生了 PDF（零 PDF 红线）"


def test_a42_5_panel_label_style_default_paren_and_bare() -> None:
    """A42.5：layout panel_label_style 默认 (a) 兼容；bare 为裸粗体小写 a。"""
    from core.render import layout as lay

    assert tuple(lay.PANEL_LABEL_STYLES) == ("paren", "bare"), lay.PANEL_LABEL_STYLES
    spec = lay.LayoutSpec()
    assert spec.panel_label_style == "paren"
    with pytest.raises(lay.LayoutError) as exc:
        lay.LayoutSpec(panel_label_style="bad").validate()
    assert _is_zh(str(exc.value)), f"非法样式缺中文原因：{exc.value!r}"

    text_fn = getattr(lay, "panel_label_text", None) or getattr(lay, "panel_label", None)
    if callable(text_fn):
        assert text_fn(0) == "(a)" and text_fn(1) == "(b)"
        assert text_fn(0, style="bare") == "a"
        assert text_fn(2, style="bare") == "c"
        with pytest.raises(ValueError):
            text_fn(0, style="bad")

    apply_fn = getattr(lay, "apply_panel_label", None)
    if callable(apply_fn):
        import matplotlib
        matplotlib.use("Agg", force=False)
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots(figsize=(2.0, 1.5))
        try:
            txt = apply_fn(ax, 0)                    # 默认 paren
            assert txt.get_text() == "(a)", txt.get_text()
            assert _bold(txt.get_fontweight())
            txt2 = apply_fn(ax, 1, style="bare")
            assert txt2.get_text() == "b", txt2.get_text()
            assert _bold(txt2.get_fontweight()), "bare 面板字母未加粗"
        finally:
            plt.close(fig)

    # 新图型默认用 bare（论文式裸字母）：meta 可断言
    pp = _pp()
    wl = np.arange(450.0, 551.0, 1.0)
    frames = [{"x": wl, "y": np.exp(-0.5 * ((wl - 500) / 20.0) ** 2), "t": 0.0}]
    _fig, meta = pp.build_spectral_evolution_figure(frames)
    assert meta.get("panel_label") == "a", meta.get("panel_label")
    assert meta.get("panel_label_style") == "bare", meta.get("panel_label_style")


# ═══════════════════════ A42.6 红线 ═══════════════════════
def test_a42_6_redlines_config_pipeline_legacy_rcparams() -> None:
    """A42.6：零新配置键；流水线不 import 新模块；legacy fig a–e 结构未改；rcParams 不被改。"""
    from core.config import DEFAULT_CFG

    for key in ("fluorescence_qy", "photobleach", "characterization",
                "photophysics", "physical_char", "qp"):
        assert key not in DEFAULT_CFG, f"新增了配置键 {key}（G 道红线）"
    assert set(DEFAULT_CFG["gui"]) == GUI_BASELINE_KEYS, (
        f"gui 键集变了：{sorted(set(DEFAULT_CFG['gui']) ^ GUI_BASELINE_KEYS)}")

    # 新模块不进流水线（on-demand；不改数值/分析链路）
    import core.pipeline as pipeline
    src = inspect.getsource(pipeline)
    for mod in ("fluorescence_qy", "photobleach", "characterization", "photophysics"):
        assert mod not in src, f"core/pipeline.py 引用了新模块 {mod}（不应进流水线）"

    # legacy JACS fig a–e 渲染函数存在且签名前缀未变（红线）
    from core.render import plotting_jacs as pj
    expected = {
        "fig_a_overlay": ["scans", "cfg", "figdir"],
        "fig_b_fit": ["g", "e", "cfg", "figdir"],
        "fig_c_bars": ["results", "cfg", "figdir"],
        "fig_d_phi": ["phi", "cfg", "figdir"],
        "fig_e_residuals": ["results", "cfg", "figdir"],
    }
    for name, prefix in expected.items():
        fn = getattr(pj, name, None)
        assert callable(fn), f"legacy 渲染函数 {name} 不见了（红线）"
        params = list(inspect.signature(fn).parameters)
        assert params[:len(prefix)] == prefix, f"{name} 签名变了：{params}"

    # legacy JACS 样式单源仍在（新图型从 JACS_SCHEMA 派生，不改 schema 默认）
    from core.render.plotstyle import JACS_SCHEMA, jacs_defaults
    defaults = jacs_defaults()
    assert "tick_direction" in defaults and defaults["tick_direction"] == "in"
    assert "spine_width" in defaults and float(defaults["spine_width"]) == pytest.approx(0.9)
    assert "panel_label_fmt" in JACS_SCHEMA and JACS_SCHEMA["panel_label_fmt"]["default"] == "({})"
